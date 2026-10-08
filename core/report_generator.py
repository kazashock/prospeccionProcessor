"""Generacion de los Excel de salida (normalizada, descartada) con formato."""

from io import BytesIO

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

HEADER_FILL = PatternFill(start_color="FF1F4E78", end_color="FF1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFFFF", bold=True)


def _style_sheet(worksheet, df: pd.DataFrame) -> None:
    for col_idx, column in enumerate(df.columns, start=1):
        cell = worksheet.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")

        max_len = max([len(str(column))] + [len(str(v)) for v in df[column].head(200)])
        worksheet.column_dimensions[cell.column_letter].width = min(max(max_len + 2, 10), 40)

    worksheet.freeze_panes = "A2"
    if worksheet.max_row > 1:
        worksheet.auto_filter.ref = worksheet.dimensions


def build_normalizada_workbook(df_validos: pd.DataFrame, audit_df: pd.DataFrame) -> BytesIO:
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        df_validos.to_excel(writer, sheet_name="Prospeccion_Normalizada", index=False)
        audit_df.to_excel(writer, sheet_name="Auditoria", index=False)

        _style_sheet(writer.sheets["Prospeccion_Normalizada"], df_validos)
        _style_sheet(writer.sheets["Auditoria"], audit_df)

    buffer.seek(0)
    return buffer


def build_descartada_workbook(df_descartados: pd.DataFrame) -> BytesIO:
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        df_descartados.to_excel(writer, sheet_name="Prospeccion_Descartada", index=False)
        _style_sheet(writer.sheets["Prospeccion_Descartada"], df_descartados)

    buffer.seek(0)
    return buffer
